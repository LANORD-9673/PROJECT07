"""Excel download serializer for the Python app (no server-side file write)."""
import io
import math

XLSX_MIME = 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'


def table_xlsx(frame, sheet_name='Du lieu'):
    """Each source field is a cell; strings stay literal, including leading zeros."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    book = Workbook()
    sheet = book.active
    sheet.title = sheet_name
    sheet.append([str(c) for c in frame.columns])
    for record in frame.itertuples(index=False, name=None):
        sheet.append(list(record))
    widths = {}
    for row in sheet:
        for cell in row:
            if isinstance(cell.value, str):
                # The value is data, not a formula, even when it begins with '='.
                cell.data_type = 's'
                cell.number_format = '@'
            value = '' if cell.value is None else str(cell.value)
            widths[cell.column] = max(widths.get(cell.column, 14),
                                      min(60, max((len(s) for s in value.splitlines()), default=0) + 2))
            cell.alignment = Alignment(vertical='top', wrap_text=True)
            if cell.row == 1:
                cell.font = Font(name='Calibri', size=11, bold=True, color='FFFFFF')
                cell.fill = PatternFill('solid', fgColor='16324F')
            else:
                cell.font = Font(name='Calibri', size=11)
                if cell.row % 2 == 0:
                    cell.fill = PatternFill('solid', fgColor='F1F6FB')
    for column, width in widths.items():
        sheet.column_dimensions[get_column_letter(column)].width = width
    for row in sheet:
        lines = max((sum(max(1, math.ceil(len(s) / max(1, widths[c.column] - 3)))
                         for s in (str(c.value or '').splitlines() or [''])) for c in row), default=1)
        sheet.row_dimensions[row[0].row].height = min(409, max(30 if row[0].row == 1 else 22, lines * 16))
    sheet.freeze_panes = 'A2'
    sheet.auto_filter.ref = sheet.dimensions
    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue()


def mapping_xlsx(raw):
    """Use the existing Vietnamese labels, preserving exact mapping values."""
    import csv
    import pandas as pd
    from nexus_reports import readable_mapping

    reader = csv.reader(io.StringIO(readable_mapping(raw).decode('utf-8-sig')))
    headers = next(reader)
    return table_xlsx(pd.DataFrame(list(reader), columns=headers, dtype=str), 'Doi chieu auditor')
