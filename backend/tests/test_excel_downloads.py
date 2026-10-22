"""App-runtime Excel download checks: rectangular cells and literal identifiers."""
import csv
import io
import sys
from pathlib import Path

import pandas as pd
import pytest
from openpyxl import load_workbook

sys.path.append(str(Path(__file__).resolve().parents[2]))
from nexus_exports import table_xlsx, mapping_xlsx
from services.privacy import read_table


@pytest.mark.parametrize('delimiter', [',', ';', '\t'])
def test_csv_fields_are_distinct_excel_cells_without_value_conversion(delimiter):
    columns = ['row_id', 'full_name', 'phone', 'amount_vnd', 'note']
    values = [['NX_001', 'Nguyễn An, Bình', '0901234567', '720000000', 'Dòng 1\nDòng 2'],
              ['NX_002', 'Trần Lan', '001234567890123456', '500000000', '']]
    stream = io.StringIO(newline='')
    writer = csv.writer(stream, delimiter=delimiter)
    writer.writerow(columns)
    writer.writerows(values)
    frame = read_table(stream.getvalue().encode('utf-8-sig'), 'source.csv')
    payload = table_xlsx(frame)
    sheet = load_workbook(io.BytesIO(payload)).active
    assert sheet.max_column == len(columns)
    assert sheet.max_row == len(values) + 1
    assert [c.value for c in sheet[1]] == columns
    assert sheet['B2'].value == 'Nguyễn An, Bình'
    assert sheet['C2'].value == '0901234567'
    assert sheet['C3'].value == '001234567890123456'
    assert sheet['E2'].value == 'Dòng 1\nDòng 2'
    assert read_table(payload, 'protected.xlsx').equals(frame)
    assert sheet.freeze_panes == 'A2'
    assert sheet.auto_filter.ref == 'A1:E3'


def test_excel_keeps_formula_looking_data_literal_and_typed_numbers_numeric():
    frame = pd.DataFrame([['=HYPERLINK("https://example.invalid")', '001', 720000000]],
                         columns=['=header', 'customer_id', 'amount_vnd'])
    sheet = load_workbook(io.BytesIO(table_xlsx(frame))).active
    assert sheet['A1'].data_type == sheet['A2'].data_type == 's'
    assert sheet['A2'].value == frame.iloc[0, 0]
    assert sheet['B2'].value == '001'
    assert sheet['C2'].value == 720000000 and sheet['C2'].data_type == 'n'


def test_auditor_mapping_has_separate_vietnamese_columns_and_exact_values():
    stream = io.StringIO(newline='')
    writer = csv.writer(stream)
    writer.writerow(['data_row', 'column', 'action', 'original_row_id', 'protected_row_id',
                     'method', 'original_value', 'protected_value'])
    writer.writerow(['1', 'full_name', 'TOKENIZE', '001', 'NX_01',
                     'HMAC-SHA256 scope token', 'Nguyễn An, Bình', 'NX_name'])
    sheet = load_workbook(io.BytesIO(mapping_xlsx(stream.getvalue().encode('utf-8-sig')))).active
    assert (sheet.max_row, sheet.max_column) == (2, 8)
    assert sheet['G1'].value == 'Giá trị gốc'
    assert sheet['D2'].value == '001'
    assert sheet['G2'].value == 'Nguyễn An, Bình'
    assert sheet['H2'].value == 'NX_name'
