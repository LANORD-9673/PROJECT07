"""Readable audit views and Word exports. No AI calls, keys or invented approvals."""
import io
import json
import csv
from datetime import datetime
from zoneinfo import ZoneInfo

VN = ZoneInfo('Asia/Ho_Chi_Minh')
STATUSES = {
    'PASS': 'Đạt yêu cầu', 'EXCEPTION': 'Phát hiện ngoại lệ',
    'INSUFFICIENT_INFORMATION': 'Chưa đủ thông tin', 'BLOCKED': 'Chưa thể kiểm tra',
    'OUT_OF_SCOPE': 'Ngoài phạm vi', 'NOT_APPLICABLE': 'Không áp dụng',
    'PENDING': 'Chờ review', 'CONFIRMED': 'Đã xác nhận',
    'NEEDS_FOLLOWUP': 'Cần làm rõ thêm', 'REJECTED': 'Không chấp nhận',
    'APPROVED': 'Đã duyệt', 'DRAFT': 'Bản nháp', 'CREATED': 'Đã khởi tạo',
    'SUPPORTED': 'Có thể kiểm tra tự động', 'UNSUPPORTED_REQUIREMENT': 'Cần thủ tục bổ sung',
    'LIVE_SUCCESS': 'Google trả kết quả thành công', 'ERROR': 'Lời gọi không thành công',
    'TEST_STUB_SUCCESS': 'Kết quả giả lập kiểm thử',
}
TESTS = {'T01': 'T01 · Phân tách người lập và người duyệt',
         'T02': 'T02 · Hạn mức phê duyệt', 'T03': 'T03 · Phê duyệt trước giải ngân'}
SOURCES = {'policy': 'Chính sách và quy trình', 'policies': 'Chính sách và quy trình',
           'loans': 'Hồ sơ khoản vay', 'authority': 'Bảng phân quyền phê duyệt', 'authorities': 'Bảng phân quyền phê duyệt'}
KINDS = {'RCM': 'Ma trận rủi ro và kiểm soát', 'PLAN': 'Kế hoạch kiểm tra',
         'FINDINGS': 'Phát hiện kiểm toán', 'REPORT': 'Báo cáo kiểm toán'}
TASKS = {'LLM1_RCM': 'Lập ma trận rủi ro và kiểm soát', 'LLM2_PLAN': 'Lập kế hoạch kiểm tra',
         'LLM3_FINDINGS': 'Soạn phát hiện kiểm toán', 'LLM4_REPORT': 'Soạn báo cáo kiểm toán'}
ACTIONS = {'TOKENIZE': 'Thay bằng mã giả', 'MASK': 'Che giá trị', 'KEEP': 'Giữ nguyên'}
FIELDS = {
    'row_id': 'Mã dòng hồ sơ', 'loan_id': 'Mã khoản vay', 'customer_id': 'Mã khách hàng',
    'full_name': 'Họ và tên', 'customer_name': 'Tên khách hàng', 'phone': 'Số điện thoại',
    'email': 'Email', 'address': 'Địa chỉ', 'maker_id': 'Mã người lập',
    'approver_id': 'Mã người phê duyệt', 'employee_id': 'Mã nhân sự', 'branch_id': 'Mã chi nhánh',
    'amount_vnd': 'Số tiền vay (VND)', 'currency': 'Loại tiền', 'role': 'Chức danh',
    'authority_limit_vnd': 'Hạn mức phê duyệt (VND)', 'approval_status': 'Trạng thái phê duyệt',
    'approval_time': 'Thời điểm phê duyệt', 'disbursement_time': 'Thời điểm giải ngân',
    'disbursement_status': 'Trạng thái giải ngân', 'valid_from': 'Hiệu lực từ', 'valid_to': 'Hiệu lực đến',
    'id_document': 'Số giấy tờ nhận dạng', 'loan_product': 'Sản phẩm tín dụng', 'loan_purpose': 'Mục đích vay',
    'term_months': 'Kỳ hạn vay (tháng)', 'interest_rate_pct': 'Lãi suất (%/năm)', 'collateral_value_vnd': 'Giá trị tài sản bảo đảm (VND)',
}
REASONS = {
    'SEGREGATION_OK': 'Người lập và người phê duyệt khác nhau.',
    'MAKER_EQUALS_APPROVER': 'Người lập đồng thời là người phê duyệt.',
    'MISSING_MAKER_OR_APPROVER': 'Thiếu mã người lập hoặc người phê duyệt.',
    'NOT_APPROVED': 'Hồ sơ chưa ở trạng thái đã phê duyệt nên không áp dụng thủ tục này.',
    'MISSING_APPROVAL_STATUS': 'Thiếu trạng thái phê duyệt.',
    'MISSING_AUTHORITY_INPUT': 'Thiếu số tiền, loại tiền, người duyệt hoặc thời điểm phê duyệt.',
    'NON_VND': 'Khoản vay không bằng VND, ngoài phạm vi bộ quy tắc hiện tại.',
    'AMOUNT_ABOVE_2B': 'Khoản vay trên 2 tỷ đồng, ngoài phạm vi bộ quy tắc hiện tại.',
    'MISSING_BRANCH': 'Thiếu thông tin chi nhánh.',
    'INVALID_APPROVAL_TIME': 'Thời điểm phê duyệt không hợp lệ.',
    'INVALID_AUTHORITY_RECORD': 'Bảng phân quyền có bản ghi không hợp lệ.',
    'NO_ACTIVE_AUTHORITY': 'Không tìm thấy phân quyền còn hiệu lực tại thời điểm phê duyệt.',
    'OVERLAPPING_AUTHORITY_CONFLICT': 'Có các phân quyền chồng lấn với hạn mức hoặc chức danh mâu thuẫn.',
    'MISSING_ROLE': 'Thiếu chức danh của người phê duyệt.',
    'POLICY_AUTHORITY_NOT_VERIFIED': 'Bảng phân quyền chưa được xác nhận theo chính sách áp dụng.',
    'MISSING_DISBURSEMENT_STATUS': 'Thiếu trạng thái giải ngân.',
    'UNKNOWN_DISBURSEMENT_STATUS': 'Trạng thái giải ngân không thuộc các giá trị được hỗ trợ.',
    'NOT_DISBURSED': 'Khoản vay chưa giải ngân nên không áp dụng thủ tục này.',
    'DISBURSED_WHILE_REJECTED': 'Khoản vay đã giải ngân dù bị từ chối phê duyệt.',
    'DISBURSED_WITHOUT_APPROVAL': 'Khoản vay đã giải ngân khi chưa có phê duyệt hợp lệ.',
    'MISSING_OR_INVALID_TIMESTAMP': 'Thiếu hoặc sai thời điểm phê duyệt hay giải ngân.',
    'APPROVAL_NOT_BEFORE_DISBURSEMENT': 'Phê duyệt xảy ra cùng lúc hoặc sau thời điểm giải ngân.',
    'APPROVED_BEFORE_DISBURSEMENT': 'Phê duyệt được thực hiện trước khi giải ngân.',
}
for tier in ('TIER_1', 'TIER_2'):
    REASONS['WITHIN_APPROVAL_LIMIT_' + tier] = 'Số tiền nằm trong hạn mức của người phê duyệt.'
    REASONS['APPROVAL_LIMIT_EXCEEDED_' + tier] = 'Số tiền vượt hạn mức của người phê duyệt.'


def label(value):
    return STATUSES.get(value, str(value or 'Chưa có thông tin'))


def field(value):
    return FIELDS.get(value, str(value).replace('_', ' '))


def reason(value):
    return REASONS.get(value, 'Cần auditor đọc và đối chiếu nguyên nhân: ' + str(value or 'chưa được ghi nhận'))


def when(value):
    if not value: return 'Chưa ghi nhận'
    try:
        d = datetime.fromisoformat(str(value).replace('Z', '+00:00'))
        if d.tzinfo is not None: d = d.astimezone(VN)
        return d.strftime('%d/%m/%Y %H:%M')
    except ValueError: return str(value)


def plain(value):
    if value is None: return 'Chưa có thông tin'
    if isinstance(value, bool): return 'Có' if value else 'Không'
    if isinstance(value, list): return '; '.join(plain(x) for x in value) or 'Không có'
    if isinstance(value, dict): return '\n'.join(f'{field(k)}: {plain(v)}' for k, v in value.items()) or 'Không có'
    return str(value)


def control_rows(items):
    return [{'Kiểm soát': x.get('control_id'), 'Rủi ro': x.get('risk'), 'Yêu cầu kiểm soát': x.get('control'),
             'Điều khoản': x.get('clause'), 'Trích dẫn chính sách': x.get('source_quote'),
             'Thủ tục kiểm tra': '; '.join(TESTS.get(t, t) for t in x.get('test_ids', []))} for x in items]


def plan_rows(items):
    return [{'Thủ tục': TESTS.get(x.get('test_id'), x.get('test_id')), 'Kiểm soát': x.get('control_id'),
             'Mục đích': x.get('purpose'), 'Bằng chứng cần thu thập': x.get('evidence_required'),
             'Điều kiện áp dụng': x.get('applicability'), 'Khả năng thực hiện': label(x.get('support'))} for x in items]


def result_rows(items):
    return [{'Hồ sơ': x.get('row_id'), 'Thủ tục': TESTS.get(x.get('test_id'), x.get('test_id')),
             'Kết quả': label(x.get('test_result')), 'Diễn giải': reason(x.get('reason_code')),
             'Review của auditor': label(x.get('human_review', 'PENDING'))} for x in items]


def evidence_rows(item):
    return [{'Nguồn bằng chứng': SOURCES.get(x.get('source'), x.get('source')), 'Dòng trong file': x.get('source_line'),
             'Hồ sơ': x.get('row_id') or item.get('row_id') or 'Không áp dụng',
             'Trường đối chiếu': ', '.join(field(k) for k in x.get('fields', [])) or 'Bản ghi phân quyền tương ứng',
             'Mã bằng chứng': x.get('evidence_id') or 'Chưa ghi nhận'} for x in item.get('evidence_refs', [])]


def snapshot_rows(items):
    return [{'Tên file': x.get('filename'), 'Loại tài liệu': SOURCES.get(x.get('source_type'), x.get('source_type')),
             'Phiên bản': x.get('version'), 'Ngày tiếp nhận': when(x.get('imported_at')),
             'Tình trạng sử dụng': 'Đang gắn cuộc kiểm toán' if x.get('used_by_runs') else 'Có thể xóa'} for x in items]


def call_rows(items):
    return [{'Thời điểm': when(x.get('created_at')), 'Tác vụ': TASKS.get(x.get('task'), x.get('task')),
             'Model đã dùng': x.get('model'), 'Thời gian (giây)': round(x.get('latency_ms', 0) / 1000, 2),
             'Kết quả gọi AI': label(x.get('status'))} for x in items]


def event_rows(items):
    rows = []
    actions = {'AUDIT_RUN_CREATED': 'Khởi tạo cuộc kiểm toán', 'P2_TESTS_EXECUTED': 'Chạy các thủ tục kiểm tra dữ liệu',
               'RESULT_REVIEWED': 'Review kết quả kiểm tra', 'FINDING_REVIEWED': 'Review phát hiện kiểm toán'}
    for x in items:
        try: d = json.loads(x.get('details_json', '{}'))
        except (ValueError, TypeError): d = {}
        action = x.get('action', '')
        name = actions.get(action)
        if name is None:
            name = next((KINDS[k] + (' được duyệt' if action.endswith('_APPROVED') else ' được tạo')
                         for k in KINDS if action in (k + '_APPROVED', k + '_GENERATED')), action.replace('_', ' ').capitalize())
        details = []
        if d.get('reason'): details.append(d['reason'])
        if d.get('results') is not None: details.append(str(d['results']) + ' lượt kiểm tra')
        if d.get('model'): details.append('Model: ' + str(d['model']))
        if d.get('finding_key'): details.append('Phát hiện ' + str(d['finding_key']))
        state = d.get('review_state') or d.get('state')
        if state: details.append(label(state))
        rows.append({'Thời điểm': when(x.get('created_at')), 'Hoạt động': name,
                     'Người thực hiện review': d.get('reviewer') or 'Hệ thống / chưa ghi nhận danh tính',
                     'Nội dung ghi nhận': '; '.join(details) or 'Đã ghi nhận hoạt động và liên kết tài liệu nguồn.'})
    return rows


def table_section(title, rows, text=None):
    return {'title': title, 'rows': rows, 'text': text or ''}


def document_sections(document):
    payload = document.get('payload', {})
    sections = []
    if document.get('kind') == 'RCM':
        sections.append(table_section('Ma trận rủi ro và kiểm soát', control_rows(payload.get('controls', []))))
    elif document.get('kind') == 'PLAN':
        sections.append(table_section('Kế hoạch kiểm tra', plan_rows(payload.get('tests', []))))
    elif document.get('kind') in ('FINDINGS', 'REPORT'):
        for key, title in {'executive_summary': 'Tóm tắt điều hành', 'scope': 'Phạm vi kiểm toán',
                           'methods': 'Phương pháp thực hiện', 'findings_summary': 'Tổng hợp phát hiện',
                           'limitations': 'Giới hạn của cuộc kiểm toán', 'recommendations': 'Kiến nghị'}.items():
            if isinstance(payload.get(key), str): sections.append(table_section(title, [], payload[key]))
        for item in payload.get('findings', payload.get('evidence_appendix', [])):
            text = '\n'.join(f'{title}: {plain(item.get(k))}' for k, title in {
                'criteria': 'Tiêu chí', 'condition': 'Thực trạng', 'risk': 'Rủi ro', 'recommendation': 'Kiến nghị',
                'source_quote': 'Trích dẫn chính sách'}.items())
            sections.append(table_section(str(item.get('finding_key', '')) + ' ' + str(item.get('title', 'Phát hiện')), [], text))
            for evidence in item.get('evidence_details', []):
                rows = [{'Trường dữ liệu': field(k), 'Giá trị được quan sát': plain(v)} for k, v in evidence.get('observed_values', {}).items()]
                sections.append(table_section('Bằng chứng hồ sơ ' + str(evidence.get('row_id', '')), rows, reason(evidence.get('reason_code'))))
    if isinstance(payload.get('limitations'), list) and payload['limitations']:
        sections.append(table_section('Nội dung cần lưu ý', [], '\n'.join(str(x) for x in payload['limitations'])))
    return sections


def evaluation_report(evaluation, wf, events, results, sources=None):
    run = wf['run']; docs = {x['kind']: x for x in wf.get('documents', [])}
    reviews = {x['result_id']: x['review_state'] for x in wf.get('result_reviews', [])}
    exceptions = [x for x in results if x.get('test_result') == 'EXCEPTION']
    pending = sum(reviews.get(x['result_id']) not in ('CONFIRMED', 'REJECTED') for x in exceptions)
    confirmed = sum(reviews.get(x['result_id']) == 'CONFIRMED' for x in exceptions)
    finding_states = {x['finding_key']: x['state'] for x in wf.get('finding_reviews', [])}
    findings = docs.get('FINDINGS', {}).get('payload', {}).get('findings', [])
    findings_pending = sum(finding_states.get(x['finding_key']) != 'CONFIRMED' for x in findings)
    unresolved = sum(evaluation.get('counts', {}).get(k, 0) for k in ('INSUFFICIENT_INFORMATION', 'BLOCKED'))
    incomplete = [KINDS[k] for k in ('RCM', 'PLAN', 'REPORT') if docs.get(k, {}).get('state') != 'APPROVED']
    if not results: incomplete.append('Kiểm tra dữ liệu')
    if pending: incomplete.append('Review ngoại lệ')
    if unresolved: incomplete.append('Làm rõ dữ liệu chưa đủ thông tin hoặc bị chặn')
    if not findings or findings_pending: incomplete.append('Review phát hiện kiểm toán')
    conclusion = ('Còn nội dung cần hoàn tất: ' + '; '.join(incomplete) + '.') if incomplete else 'Các bước review đã được ghi nhận. Auditor cần đối chiếu giới hạn trước khi sử dụng báo cáo; ứng dụng chưa cho phép phát hành chính thức.'
    totals = evaluation.get('counts', {})
    summary = f"Đã ghi nhận {evaluation.get('total', 0)} lượt kiểm tra trên {len({x.get('row_id') for x in results})} hồ sơ, gồm {len(exceptions)} lượt phát hiện ngoại lệ. Trong số ngoại lệ, {confirmed} đã được auditor xác nhận và {pending} còn chờ xử lý. Một hồ sơ có thể phát sinh nhiều lượt kiểm tra hoặc ngoại lệ."
    stages = [{'Hạng mục': KINDS[k], 'Trạng thái': label(docs[k]['state']) if k in docs else 'Chưa thực hiện',
               'Người review': docs.get(k, {}).get('reviewed_by') or 'Chưa ghi nhận',
               'Thời điểm review': when(docs.get(k, {}).get('reviewed_at'))} for k in KINDS]
    if findings:
        row = stages[2]
        row['Trạng thái'] = f'{len(findings) - findings_pending}/{len(findings)} phát hiện đã được xác nhận'
        recorded = [x for x in wf.get('finding_reviews', []) if x.get('finding_key') in {f['finding_key'] for f in findings}]
        row['Người review'] = ', '.join(sorted({x['reviewer'] for x in recorded if x.get('reviewer')})) or 'Chưa ghi nhận'
        row['Thời điểm review'] = when(max((x.get('reviewed_at', '') for x in recorded), default=''))
    metrics = []
    descriptions = {
        'quote_match_rate': ('Trích dẫn RCM có trong chính sách', 'Đối chiếu nguyên văn; chưa đánh giá đúng nghĩa hoặc đủ nội dung.'),
        'test_coverage': ('Độ phủ ba thủ tục T01 đến T03', 'Tỷ lệ thủ tục có hỗ trợ thực thi trong kế hoạch; không đại diện mọi kiểm soát.'),
        'evidence_link_integrity': ('Liên kết bằng chứng hợp lệ', 'Kiểm tra mã tham chiếu tồn tại trong cuộc kiểm toán; không thay cho review bằng chứng.'),
    }
    for key, (title, note) in descriptions.items():
        value = evaluation.get('structural_metrics', {}).get(key)
        metrics.append({'Chỉ tiêu': title, 'Kết quả': f'{value:.1%}' if value is not None else 'Chưa có dữ liệu',
                        'Nhận xét': 'Đủ theo phép kiểm tra cấu trúc' if value == 1 else 'Cần rà soát' if value is not None else 'Chưa đánh giá', 'Cơ sở và giới hạn': note})
    benchmark = evaluation.get('regression_benchmark', {})
    regression = [{'Chỉ tiêu đối chứng': title, 'Kết quả': f'{benchmark[key]:.1%}' if benchmark.get(key) is not None else 'Chưa tính',
                   'Phạm vi': str(benchmark.get('n', 0)) + ' tình huống quy tắc cố định'} for key, title in {
                       'status_accuracy': 'Kết quả khớp đáp án', 'reason_accuracy': 'Kết quả và nguyên nhân cùng khớp',
                       'exception_precision': 'Tỷ lệ ngoại lệ đúng trong số ngoại lệ phát hiện',
                       'exception_recall': 'Tỷ lệ tìm được ngoại lệ đã cài sẵn'}.items()]
    counts = [{'Kết quả kiểm tra': label(k), 'Số lượt': totals.get(k, 0)} for k in (
        'PASS', 'EXCEPTION', 'INSUFFICIENT_INFORMATION', 'BLOCKED', 'OUT_OF_SCOPE', 'NOT_APPLICABLE')]
    followups = []
    if pending: followups.append({'Việc cần hoàn tất': 'Review ' + str(pending) + ' ngoại lệ chưa kết luận', 'Người xử lý': 'Auditor được phân công', 'Trạng thái': 'Chờ xử lý'})
    if unresolved: followups.append({'Việc cần hoàn tất': 'Làm rõ dữ liệu cho ' + str(unresolved) + ' lượt chưa đủ thông tin hoặc chưa thể kiểm tra', 'Người xử lý': 'Auditor và đơn vị cung cấp dữ liệu', 'Trạng thái': 'Cần bổ sung bằng chứng'})
    for k in ('RCM', 'PLAN', 'REPORT'):
        if docs.get(k, {}).get('state') != 'APPROVED': followups.append({'Việc cần hoàn tất': 'Hoàn tất và review ' + KINDS[k].lower(), 'Người xử lý': 'Auditor được phân công', 'Trạng thái': 'Chưa hoàn tất'})
    if not findings or findings_pending: followups.append({'Việc cần hoàn tất': 'Hoàn tất và review các phát hiện kiểm toán', 'Người xử lý': 'Auditor được phân công', 'Trạng thái': 'Chưa hoàn tất'})
    limitations = 'Các tỷ lệ cấu trúc chỉ kiểm tra trích dẫn, độ phủ và liên kết; chưa chứng minh AI diễn giải đúng nghiệp vụ. Độ chính xác, precision, recall và mức bám sát nguồn của AI trên hồ sơ cuộc kiểm toán chưa được đo vì chưa có đáp án độc lập đã xác nhận. Bộ đối chứng quy tắc do nhóm tự xây dựng, chưa được chứng nhận độc lập. Danh tính reviewer trong demo do người dùng tự khai. Báo cáo vẫn cần auditor review và chưa được phát hành chính thức.'
    sections = [table_section('Kết luận tổng hợp', [], conclusion + '\n' + summary),
                table_section('Tình trạng thực hiện và review', stages), table_section('Kết quả kiểm tra dữ liệu', counts),
                table_section('Chỉ tiêu cấu trúc và khả năng truy vết', metrics),
                table_section('Công việc cần tiếp tục', followups, '' if followups else 'Chưa ghi nhận công việc tồn đọng theo các tiêu chí hoàn tất ở trên. Các giới hạn đánh giá vẫn áp dụng.'),
                table_section('Đối chứng bộ quy tắc phần mềm', regression, 'Kết quả dưới đây thuộc bộ kiểm thử cố định của Python, tách biệt với độ chính xác của AI trên cuộc kiểm toán hiện tại.'),
                table_section('Giới hạn và lưu ý sử dụng', [], limitations)]
    if sources: sections.append(table_section('Phụ lục tài liệu nguồn', [{'Tài liệu': x.get('filename'), 'Loại': SOURCES.get(x.get('source_type')), 'Phiên bản': x.get('version')} for x in sources]))
    cases = [{'Tình huống': x['case_id'], 'Thủ tục': x['test_id'], 'Kỳ vọng': label(x['expected']),
              'Thực tế': label(x['actual']), 'Đối chiếu': 'Khớp' if x['matched'] else 'Cần rà soát'} for x in benchmark.get('cases', [])]
    sections.extend([table_section('Phụ lục kiểm thử đối chứng', cases), table_section('Phụ lục lịch sử gọi AI', call_rows(evaluation.get('llm_calls', []))), table_section('Phụ lục nhật ký kiểm toán', event_rows(events))])
    sections[-3]['text'] = 'Đối chiếu được ghi là khớp khi cả kết quả và nguyên nhân cùng đúng với đáp án của tình huống.'
    sections[-2]['text'] = 'Thời gian gồm cả lời gọi thất bại. Google trả kết quả thành công chưa chứng minh nội dung đã đạt kiểm tra nghiệp vụ hoặc được auditor duyệt.'
    return {'title': 'Báo cáo đánh giá cuộc kiểm toán', 'metadata': [
        ('Cuộc kiểm toán', run.get('audit_name')), ('Mã cuộc kiểm toán', run.get('run_id')),
        ('Mã báo cáo', 'NEXUS-EVAL-' + run.get('run_id', '')[:8].upper()),
        ('Thời điểm lập', datetime.now(VN).strftime('%d/%m/%Y %H:%M') + ' (giờ Việt Nam)'),
        ('Trạng thái', 'Bản nháp cần auditor review')], 'sections': sections}


def privacy_report(reviewed):
    data = json.loads(reviewed['audit_report_json_bytes'])
    rows = [{'Tên cột': x['column'], 'Cách bảo vệ': ACTIONS.get(x['action'], x['action']), 'Giải thích': x.get('reason', '')} for x in data['plan']]
    return {'title': 'Biên bản bảo vệ dữ liệu', 'metadata': [('File nguồn', data['source_filename']),
        ('Người review', data['reviewer']), ('Thời điểm review', when(data['reviewed_at'])),
        ('Phạm vi kiểm toán', data['scope']), ('Số dòng dữ liệu', data['row_count']), ('Số cột', data['column_count']),
        ('Mã lưu trữ', data['file_id']), ('Mã lần review', data['review_id'])],
        'sections': [table_section('Phương án bảo vệ dữ liệu đã duyệt', rows),
                     table_section('Lưu trữ hồ sơ đối chiếu', [], 'Bảng giá trị gốc và giá trị đã bảo vệ được tải riêng cho auditor. Biên bản này không chứa giá trị gốc từng ô. Không ghi bảng đối chiếu mới vào cơ sở dữ liệu hoặc nhật ký.')]}


def readable_mapping(raw):
    headers = {'data_row': 'Dòng dữ liệu', 'column': 'Tên cột', 'action': 'Cách bảo vệ',
               'original_row_id': 'Mã hồ sơ gốc', 'protected_row_id': 'Mã hồ sơ đã bảo vệ',
               'method': 'Phương pháp', 'original_value': 'Giá trị gốc', 'protected_value': 'Giá trị đã bảo vệ'}
    source = csv.DictReader(io.StringIO(raw.decode('utf-8-sig')))
    output = io.StringIO(newline='')
    writer = csv.DictWriter(output, fieldnames=[headers.get(k, k) for k in source.fieldnames])
    writer.writeheader()
    for row in source:
        row['action'] = ACTIONS.get(row.get('action'), row.get('action'))
        row['method'] = {'HMAC-SHA256 scope token': 'Mã giả nhất quán trong phạm vi kiểm toán', 'literal redaction': 'Che nội dung', 'unchanged': 'Giữ nguyên dữ liệu'}.get(row.get('method'), row.get('method'))
        writer.writerow({headers.get(k, k): v for k, v in row.items()})
    return output.getvalue().encode('utf-8-sig')


def word_bytes(report):
    from docx import Document
    from docx.shared import Inches, Pt, RGBColor
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5); section.page_height = Inches(11)
    section.top_margin = section.bottom_margin = Inches(.65)
    section.left_margin = section.right_margin = Inches(.65)
    for style in ('Normal', 'Title', 'Heading 1', 'Heading 2'):
        doc.styles[style].font.name = 'Arial'
        doc.styles[style].font.color.rgb = RGBColor(0, 0, 0)
    # Remove template paragraph rules, including the built-in blue Title border.
    for node in list(doc.styles.element.xpath('.//w:pBdr')) + list(doc.element.xpath('.//w:pBdr')):
        node.getparent().remove(node)
    doc.styles['Normal'].font.size = Pt(11)
    doc.styles['Normal'].paragraph_format.space_after = Pt(7)
    doc.styles['Title'].font.size = Pt(21)
    doc.styles['Heading 1'].font.size = Pt(13)
    doc.styles['Heading 1'].paragraph_format.space_before = Pt(13)
    doc.add_paragraph('NEXUS  |  PROJECT 07').runs[0].bold = True
    doc.add_paragraph(report['title'], 'Title')
    for key, value in report.get('metadata', []):
        p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(3)
        p.add_run(str(key) + ': ').bold = True; p.add_run(plain(value))
    for part in report.get('sections', []):
        doc.add_heading(part['title'], 1)
        for text in part.get('text', '').split('\n'):
            if text: doc.add_paragraph(text)
        rows = part.get('rows', [])
        if not rows: continue
        headers = list(rows[0])
        table = doc.add_table(rows=1, cols=len(headers))
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = False
        # Give narrative columns more room than short identifiers/statuses.
        weights = [2.4 if any(k in h for k in ('Diễn giải', 'Nội dung', 'giới hạn', 'Giải thích', 'cần', 'Yêu cầu', 'Rủi ro', 'Mục đích', 'Trích dẫn', 'quan sát')) else 1.1 for h in headers]
        widths = [7.2 * w / sum(weights) for w in weights]
        for i, (header, width) in enumerate(zip(headers, widths)):
            table.columns[i].width = Inches(width); table.rows[0].cells[i].text = header
        props = table._tbl.tblPr
        borders = OxmlElement('w:tblBorders')
        for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
            node = OxmlElement('w:' + edge); node.set(qn('w:val'), 'single'); node.set(qn('w:sz'), '4'); node.set(qn('w:color'), 'D9D9D9'); borders.append(node)
        props.append(borders)
        repeat = OxmlElement('w:tblHeader'); table.rows[0]._tr.get_or_add_trPr().append(repeat)
        for row in rows:
            cells = table.add_row().cells
            for i, key in enumerate(headers): cells[i].text = plain(row.get(key))
        for n, row in enumerate(table.rows):
            no_split = OxmlElement('w:cantSplit'); row._tr.get_or_add_trPr().append(no_split)
            for i, cell in enumerate(row.cells):
                cell.width = Inches(widths[i]); cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                shade = OxmlElement('w:shd'); shade.set(qn('w:fill'), '17365D' if n == 0 else 'F2F6FA' if n % 2 == 0 else 'FFFFFF'); cell._tc.get_or_add_tcPr().append(shade)
                margins = OxmlElement('w:tcMar')
                for side in ('top', 'left', 'bottom', 'right'):
                    m = OxmlElement('w:' + side); m.set(qn('w:w'), '85'); m.set(qn('w:type'), 'dxa'); margins.append(m)
                cell._tc.get_or_add_tcPr().append(margins)
                for p in cell.paragraphs:
                    p.paragraph_format.space_after = Pt(3); p.paragraph_format.space_before = Pt(3)
                    p.paragraph_format.line_spacing = 1.05
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER if headers[i] in ('Số lượt', 'Kết quả', 'Phiên bản', 'Thủ tục', 'Tình huống', 'Thời gian (giây)', 'Đối chiếu') else WD_ALIGN_PARAGRAPH.LEFT
                    for r in p.runs:
                        r.font.size = Pt(10); r.bold = n == 0
                        r.font.color.rgb = RGBColor(255,255,255) if n == 0 else RGBColor(0,0,0)
        doc.add_paragraph().paragraph_format.space_after = Pt(0)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer.add_run('NEXUS | Trang ')
    page_field = OxmlElement('w:fldSimple'); page_field.set(qn('w:instr'), 'PAGE'); footer._p.append(page_field)
    for run in footer.runs: run.font.size = Pt(9)
    buffer = io.BytesIO(); doc.save(buffer)
    return buffer.getvalue()
